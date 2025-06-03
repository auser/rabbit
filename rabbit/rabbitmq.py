import json
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

import httpx
import pandas as pd
from dotenv import dotenv_values, load_dotenv
from openpyxl.styles import Alignment, Font, PatternFill

from rabbit.logging import logger, setup_logging
from rabbit.utils import get_definitions_url

load_dotenv() # Initial load

values = dotenv_values()

class RabbitMQ(object):
    env_files: list[str] = []
    verbose: int = 0  # Changed to int to match logging levels
    from_host: str = "localhost"
    from_port: int = 15672
    from_username: str = values.get("RABBITMQ_FROM_USERNAME", "guest")
    from_password: str = values.get("RABBITMQ_FROM_PASSWORD", "guest")
    output: str | Path = "source_rabbitmq.json"
    to_host: str = "localhost"
    to_port: int = 5672
    to_username: str = values.get("RABBITMQ_TO_USERNAME", "guest")
    to_password: str = values.get("RABBITMQ_TO_PASSWORD", "guest")
    input: str | Path = "source_rabbitmq.json"

    def __init__(self, env_files: Optional[list[str]] = None, verbose: int = 0):
        self.env_files = env_files or []
        self.verbose = verbose
        setup_logging(verbose)  # Use the verbose level directly
        load_env_files(self.env_files)

    async def download(self, output: Optional[str | Path] = None):
        """Download the rabbitmq configuration from the source host"""
        async with httpx.AsyncClient() as client:
            headers = {
                "Content-Type": "application/json"
            }
            url = get_definitions_url(self.from_host, self.from_port)
            logger.info(f"Downloading rabbitmq configuration from {url}")
            kwargs = {"auth": (self.from_username, self.from_password), "headers": headers}
            response = await client.get(url, **kwargs)
            json_str = response.text

        if output is not None:
            with open(output, "w") as f:
                f.write(json_str)

    async def upload(self, input: Optional[str | Path] = None):
        """Upload the rabbitmq configuration to the target host"""
        if input is None:
            raise ValueError("Input file is required")

        with open(input, "r") as f:
            data = json.load(f)

        async with httpx.AsyncClient() as client:
            headers = {
                "Content-Type": "application/json"
            }
            url = get_definitions_url(self.to_host, self.to_port)
            logger.info(f"Uploading rabbitmq configuration to {url}")
            kwargs = {"auth": (self.to_username, self.to_password), "headers": headers}
            response = await client.post(url, json=data, **kwargs)
            print(response.text)
        # await self._make_request("post", (self.to_username, self.to_password), data, input)
        logger.info("Uploaded rabbitmq configuration")

    async def clone(self):
        """Clone the rabbitmq configuration from the source host to the target host"""
        logger.info("Downloading rabbitmq configuration from source host")
        temp_file = tempfile.NamedTemporaryFile(delete=False)
        await self.download(temp_file.name)
        logger.info("Uploading rabbitmq configuration to target host")
        await self.upload(temp_file.name)
        logger.info("Cloned rabbitmq configuration")

    async def _make_api_request(self, endpoint: str, host: str, port: int, 
                               username: str, password: str, ssl: bool = False) -> List[Dict]:
        """Make API request to RabbitMQ management interface"""
        protocol = 'https' if ssl else 'http'
        base_url = f"{protocol}://{host}:{port}/api"
        url = f"{base_url}/{endpoint}"
        
        logger.debug(f"Making API request to: {url}")
        
        async with httpx.AsyncClient(timeout=30.0) as client:
            try:
                response = await client.get(url, auth=(username, password))
                response.raise_for_status()
                result = response.json()
                logger.debug(f"API request successful, received {len(result) if isinstance(result, list) else 'data'}")
                return result
            except httpx.RequestError as e:
                logger.error(f"Error connecting to {host}: {e}")
                raise Exception(f"Connection error: {e}")
            except httpx.HTTPStatusError as e:
                logger.error(f"HTTP error {e.response.status_code}: {e.response.text}")
                raise Exception(f"HTTP {e.response.status_code}: {e.response.text}")

    async def _collect_server_data(self, host: str, port: int, username: str, 
                                  password: str, vhost: str = "/", ssl: bool = False) -> Dict:
        """Collect all data from a RabbitMQ server"""
        logger.info(f"Collecting data from RabbitMQ server: {host}:{port}")
        
        # URL encode vhost for API calls
        vhost_encoded = vhost.replace("/", "%2F")
        
        # Collect all data
        overview = await self._make_api_request("overview", host, port, username, password, ssl)
        queues = await self._make_api_request(f"queues/{vhost_encoded}", host, port, username, password, ssl)
        exchanges = await self._make_api_request(f"exchanges/{vhost_encoded}", host, port, username, password, ssl)
        consumers = await self._make_api_request(f"consumers/{vhost_encoded}", host, port, username, password, ssl)
        bindings = await self._make_api_request(f"bindings/{vhost_encoded}", host, port, username, password, ssl)
        connections = await self._make_api_request("connections", host, port, username, password, ssl)
        channels = await self._make_api_request("channels", host, port, username, password, ssl)
        nodes = await self._make_api_request("nodes", host, port, username, password, ssl)

        data = {
            'server_info': {
                'host': host,
                'port': port,
                'vhost': vhost,
                'timestamp': datetime.now().isoformat(),
                'overview': overview if isinstance(overview, dict) else {}
            },
            'queues': queues,
            'exchanges': exchanges,
            'consumers': consumers,
            'bindings': bindings,
            'connections': connections,
            'channels': channels,
            'nodes': nodes
        }
        
        logger.info(f"✓ Found {len(queues)} queues")
        logger.info(f"✓ Found {len(exchanges)} exchanges")
        logger.info(f"✓ Found {len(consumers)} consumers")
        logger.info(f"✓ Found {len(bindings)} bindings")
        logger.info(f"✓ Found {len(connections)} connections")
        
        return data

    def _prepare_queues_data(self, queues: List[Dict]) -> pd.DataFrame:
        """Prepare queues data for export"""
        queue_data = []
        
        for queue in queues:
            message_stats = queue.get('message_stats', {})
            
            queue_info = {
                'Queue Name': queue.get('name', ''),
                'VHost': queue.get('vhost', ''),
                'Durable': queue.get('durable', False),
                'Auto Delete': queue.get('auto_delete', False),
                'State': queue.get('state', ''),
                'Messages Total': queue.get('messages', 0),
                'Messages Ready': queue.get('messages_ready', 0),
                'Messages Unacknowledged': queue.get('messages_unacknowledged', 0),
                'Active Consumers': queue.get('consumers', 0),
                'Memory (bytes)': queue.get('memory', 0),
                'Message Bytes': queue.get('message_bytes', 0),
                'Message Rate/sec': message_stats.get('publish_details', {}).get('rate', 0),
                'Consumer Utilisation': queue.get('consumer_utilisation', ''),
                'Node': queue.get('node', ''),
                'Arguments': json.dumps(queue.get('arguments', {})) if queue.get('arguments') else '',
                'Policy': queue.get('policy', ''),
                'Exclusive': queue.get('exclusive', False),
                'Idle Since': queue.get('idle_since', ''),
            }
            queue_data.append(queue_info)
        
        return pd.DataFrame(queue_data)

    def _prepare_exchanges_data(self, exchanges: List[Dict]) -> pd.DataFrame:
        """Prepare exchanges data for export"""
        exchange_data = []
        
        for exchange in exchanges:
            message_stats = exchange.get('message_stats', {})
            
            exchange_info = {
                'Exchange Name': exchange.get('name', ''),
                'VHost': exchange.get('vhost', ''),
                'Type': exchange.get('type', ''),
                'Durable': exchange.get('durable', False),
                'Auto Delete': exchange.get('auto_delete', False),
                'Internal': exchange.get('internal', False),
                'Messages In': message_stats.get('publish_in', 0),
                'Messages Out': message_stats.get('publish_out', 0),
                'Message Rate In/sec': message_stats.get('publish_in_details', {}).get('rate', 0),
                'Message Rate Out/sec': message_stats.get('publish_out_details', {}).get('rate', 0),
                'Arguments': json.dumps(exchange.get('arguments', {})) if exchange.get('arguments') else '',
                'Policy': exchange.get('policy', ''),
            }
            exchange_data.append(exchange_info)
        
        return pd.DataFrame(exchange_data)

    def _prepare_consumers_data(self, consumers: List[Dict]) -> pd.DataFrame:
        """Prepare consumers data for export"""
        consumer_data = []
        
        for consumer in consumers:
            queue_info = consumer.get('queue', {})
            channel_info = consumer.get('channel_details', {})
            
            consumer_info = {
                'Consumer Tag': consumer.get('consumer_tag', ''),
                'Queue Name': queue_info.get('name', ''),
                'Queue VHost': queue_info.get('vhost', ''),
                'Channel': channel_info.get('name', ''),
                'Connection Name': channel_info.get('connection_name', ''),
                'Peer Host': channel_info.get('peer_host', ''),
                'Peer Port': channel_info.get('peer_port', ''),
                'User': channel_info.get('user', ''),
                'Prefetch Count': consumer.get('prefetch_count', 0),
                'Ack Required': consumer.get('ack_required', True),
                'Exclusive': consumer.get('exclusive', False),
                'Arguments': json.dumps(consumer.get('arguments', {})) if consumer.get('arguments') else '',
            }
            consumer_data.append(consumer_info)
        
        return pd.DataFrame(consumer_data)

    def _prepare_bindings_data(self, bindings: List[Dict]) -> pd.DataFrame:
        """Prepare bindings data for export"""
        binding_data = []
        
        for binding in bindings:
            binding_info = {
                'Source Exchange': binding.get('source', ''),
                'Destination Type': binding.get('destination_type', ''),
                'Destination': binding.get('destination', ''),
                'Routing Key': binding.get('routing_key', ''),
                'VHost': binding.get('vhost', ''),
                'Properties Key': binding.get('properties_key', ''),
                'Arguments': json.dumps(binding.get('arguments', {})) if binding.get('arguments') else '',
            }
            binding_data.append(binding_info)
        
        return pd.DataFrame(binding_data)

    def _prepare_connections_data(self, connections: List[Dict]) -> pd.DataFrame:
        """Prepare connections data for export"""
        connection_data = []
        
        for conn in connections:
            conn_info = {
                'Connection Name': conn.get('name', ''),
                'Node': conn.get('node', ''),
                'User': conn.get('user', ''),
                'VHost': conn.get('vhost', ''),
                'State': conn.get('state', ''),
                'Protocol': conn.get('protocol', ''),
                'Host': conn.get('host', ''),
                'Port': conn.get('port', ''),
                'Peer Host': conn.get('peer_host', ''),
                'Peer Port': conn.get('peer_port', ''),
                'SSL': conn.get('ssl', False),
                'Channels': conn.get('channels', 0),
                'Type': conn.get('type', ''),
                'Client Properties': json.dumps(conn.get('client_properties', {})) if conn.get('client_properties') else '',
                'Connected At': conn.get('connected_at', ''),
                'Timeout': conn.get('timeout', 0),
            }
            connection_data.append(conn_info)
        
        return pd.DataFrame(connection_data)

    def _prepare_summary_data(self, data: Dict) -> pd.DataFrame:
        """Prepare summary overview data"""
        overview = data['server_info']['overview']
        
        summary_info = [
            ['Server', data['server_info']['host']],
            ['Port', data['server_info']['port']],
            ['VHost', data['server_info']['vhost']],
            ['Generated At', data['server_info']['timestamp']],
            ['RabbitMQ Version', overview.get('rabbitmq_version', 'Unknown')],
            ['Erlang Version', overview.get('erlang_version', 'Unknown')],
            ['Cluster Name', overview.get('cluster_name', 'Unknown')],
            ['Total Queues', len(data['queues'])],
            ['Total Exchanges', len(data['exchanges'])],
            ['Total Consumers', len(data['consumers'])],
            ['Total Bindings', len(data['bindings'])],
            ['Total Connections', len(data['connections'])],
            ['Total Channels', len(data['channels'])],
        ]
        
        # Add message stats if available
        message_stats = overview.get('message_stats', {})
        if message_stats:
            summary_info.extend([
                ['Messages Published', message_stats.get('publish', 0)],
                ['Messages Delivered', message_stats.get('deliver_get', 0)],
                ['Messages Acknowledged', message_stats.get('ack', 0)],
            ])
        
        # Add queue totals if available
        queue_totals = overview.get('queue_totals', {})
        if queue_totals:
            summary_info.extend([
                ['Total Messages', queue_totals.get('messages', 0)],
                ['Messages Ready', queue_totals.get('messages_ready', 0)],
                ['Messages Unacknowledged', queue_totals.get('messages_unacknowledged', 0)],
            ])
        
        return pd.DataFrame(summary_info, columns=['Metric', 'Value'])

    def _format_excel_sheets(self, workbook):
        """Apply formatting to Excel sheets"""
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            
            # Format headers
            if sheet.max_row > 0:
                for cell in sheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                
                # Auto-adjust column widths
                for column in sheet.columns:
                    max_length = max(len(str(cell.value or "")) for cell in column)
                    sheet.column_dimensions[column[0].column_letter].width = min(max_length + 2, 50)

    async def inspect_server(self, host: str, port: int, username: str, password: str,
                           output: str, format: str = 'excel', vhost: str = '/', ssl: bool = False):
        """Inspect a RabbitMQ server and export data"""
        
        logger.info(f"Starting inspection of RabbitMQ server: {host}:{port}")
        logger.debug(f"Using credentials: {username}/<password hidden>")
        logger.debug(f"SSL enabled: {ssl}")
        logger.debug(f"Virtual host: {vhost}")
        
        # Collect data
        data = await self._collect_server_data(host, port, username, password, vhost, ssl)
        
        # Test connection
        if not data['server_info']['overview']:
            raise Exception("Could not connect to RabbitMQ server. Please check your connection details.")
        
        overview = data['server_info']['overview']
        logger.info(f"✓ Connected to RabbitMQ server: {host}")
        logger.info(f"  Version: {overview.get('rabbitmq_version', 'Unknown')}")
        logger.info(f"  Cluster: {overview.get('cluster_name', 'Unknown')}")
        
        # Export based on format
        if format in ['excel', 'all']:
            await self._export_to_excel(data, f"{output}.xlsx")
        
        if format in ['csv', 'all']:
            await self._export_to_csv(data, output)
        
        if format in ['json', 'all']:
            await self._export_to_json(data, f"{output}.json")
            
        logger.info("Inspection completed successfully!")

    async def _export_to_excel(self, data: Dict, filename: str):
        """Export data to Excel file with multiple sheets"""
        logger.info(f"Exporting data to Excel: {filename}")
        
        # Prepare DataFrames
        summary_df = self._prepare_summary_data(data)
        queues_df = self._prepare_queues_data(data['queues'])
        exchanges_df = self._prepare_exchanges_data(data['exchanges'])
        consumers_df = self._prepare_consumers_data(data['consumers'])
        bindings_df = self._prepare_bindings_data(data['bindings'])
        connections_df = self._prepare_connections_data(data['connections'])
        
        # Create Excel file with multiple sheets
        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            summary_df.to_excel(writer, sheet_name='Summary', index=False)
            queues_df.to_excel(writer, sheet_name='Queues', index=False)
            exchanges_df.to_excel(writer, sheet_name='Exchanges', index=False)
            consumers_df.to_excel(writer, sheet_name='Consumers', index=False)
            bindings_df.to_excel(writer, sheet_name='Bindings', index=False)
            connections_df.to_excel(writer, sheet_name='Connections', index=False)
            
            # Format the workbook
            self._format_excel_sheets(writer.book)
        
        logger.info(f"✓ Excel file created: {filename}")
        logger.info(f"  - Summary: {len(summary_df)} metrics")
        logger.info(f"  - Queues: {len(queues_df)} queues")
        logger.info(f"  - Exchanges: {len(exchanges_df)} exchanges")
        logger.info(f"  - Consumers: {len(consumers_df)} consumers")
        logger.info(f"  - Bindings: {len(bindings_df)} bindings")
        logger.info(f"  - Connections: {len(connections_df)} connections")

    async def _export_to_csv(self, data: Dict, prefix: str):
        """Export data to separate CSV files"""
        logger.info(f"Exporting data to CSV files with prefix: {prefix}")
        
        # Export each dataset to CSV
        datasets = {
            'summary': self._prepare_summary_data(data),
            'queues': self._prepare_queues_data(data['queues']),
            'exchanges': self._prepare_exchanges_data(data['exchanges']),
            'consumers': self._prepare_consumers_data(data['consumers']),
            'bindings': self._prepare_bindings_data(data['bindings']),
            'connections': self._prepare_connections_data(data['connections'])
        }
        
        for name, df in datasets.items():
            filename = f"{prefix}_{name}.csv"
            df.to_csv(filename, index=False)
            logger.info(f"✓ Created {filename} ({len(df)} rows)")

    async def _export_to_json(self, data: Dict, filename: str):
        """Export raw data to JSON file"""
        with open(filename, 'w') as f:
            json.dump(data, f, indent=2, default=str)
        logger.info(f"✓ Raw data exported to: {filename}")

    async def compare_servers(self, prod_host: str, prod_port: int, prod_username: str, prod_password: str,
                            dev_host: str, dev_port: int, dev_username: str, dev_password: str,
                            output: str, format: str = 'excel', ssl: bool = False):
        """Compare two RabbitMQ servers and export comparison data"""
        
        logger.info("Collecting production data...")
        prod_data = await self._collect_server_data(prod_host, prod_port, prod_username, prod_password, "/", ssl)
        
        logger.info("Collecting development data...")
        dev_data = await self._collect_server_data(dev_host, dev_port, dev_username, dev_password, "/", ssl)
        
        # Create comparison data structure
        comparison_data = {
            'comparison_info': {
                'generated_at': datetime.now().isoformat(),
                'production': {'host': prod_host, 'port': prod_port},
                'development': {'host': dev_host, 'port': dev_port}
            },
            'production_data': prod_data,
            'development_data': dev_data
        }
        
        # Export comparison
        if format in ['excel', 'all']:
            await self._export_comparison_to_excel(comparison_data, f"{output}.xlsx")
        
        if format in ['csv', 'all']:
            await self._export_comparison_to_csv(comparison_data, output)
        
        if format in ['json', 'all']:
            await self._export_to_json(comparison_data, f"{output}.json")
        
        # Print summary to console
        self._print_comparison_summary(comparison_data)

    async def _export_comparison_to_excel(self, comparison_data: Dict, filename: str):
        """Export comparison data to Excel with separate sheets for each environment"""
        logger.info(f"Exporting comparison to Excel: {filename}")
        
        prod_data = comparison_data['production_data']
        dev_data = comparison_data['development_data']
        
        # Prepare comparison summary
        comparison_summary = self._prepare_comparison_summary(comparison_data)
        
        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            # Comparison summary sheet
            comparison_summary.to_excel(writer, sheet_name='Comparison Summary', index=False)
            
            # Production sheets
            self._prepare_summary_data(prod_data).to_excel(writer, sheet_name='Prod Summary', index=False)
            self._prepare_queues_data(prod_data['queues']).to_excel(writer, sheet_name='Prod Queues', index=False)
            self._prepare_exchanges_data(prod_data['exchanges']).to_excel(writer, sheet_name='Prod Exchanges', index=False)
            self._prepare_consumers_data(prod_data['consumers']).to_excel(writer, sheet_name='Prod Consumers', index=False)
            self._prepare_bindings_data(prod_data['bindings']).to_excel(writer, sheet_name='Prod Bindings', index=False)
            self._prepare_connections_data(prod_data['connections']).to_excel(writer, sheet_name='Prod Connections', index=False)
            
            # Development sheets
            self._prepare_summary_data(dev_data).to_excel(writer, sheet_name='Dev Summary', index=False)
            self._prepare_queues_data(dev_data['queues']).to_excel(writer, sheet_name='Dev Queues', index=False)
            self._prepare_exchanges_data(dev_data['exchanges']).to_excel(writer, sheet_name='Dev Exchanges', index=False)
            self._prepare_consumers_data(dev_data['consumers']).to_excel(writer, sheet_name='Dev Consumers', index=False)
            self._prepare_bindings_data(dev_data['bindings']).to_excel(writer, sheet_name='Dev Bindings', index=False)
            self._prepare_connections_data(dev_data['connections']).to_excel(writer, sheet_name='Dev Connections', index=False)
            
            # Format the workbook
            self._format_excel_sheets(writer.book)
        
        logger.info(f"✓ Comparison Excel file created: {filename}")

    async def _export_comparison_to_csv(self, comparison_data: Dict, prefix: str):
        """Export comparison data to CSV files"""
        logger.info(f"Exporting comparison to CSV files with prefix: {prefix}")
        
        prod_data = comparison_data['production_data']
        dev_data = comparison_data['development_data']
        
        # Export comparison summary
        comparison_summary = self._prepare_comparison_summary(comparison_data)
        comparison_summary.to_csv(f"{prefix}_comparison_summary.csv", index=False)
        
        # Export production data
        prod_datasets = {
            'prod_summary': self._prepare_summary_data(prod_data),
            'prod_queues': self._prepare_queues_data(prod_data['queues']),
            'prod_exchanges': self._prepare_exchanges_data(prod_data['exchanges']),
            'prod_consumers': self._prepare_consumers_data(prod_data['consumers']),
            'prod_bindings': self._prepare_bindings_data(prod_data['bindings']),
            'prod_connections': self._prepare_connections_data(prod_data['connections'])
        }
        
        # Export development data
        dev_datasets = {
            'dev_summary': self._prepare_summary_data(dev_data),
            'dev_queues': self._prepare_queues_data(dev_data['queues']),
            'dev_exchanges': self._prepare_exchanges_data(dev_data['exchanges']),
            'dev_consumers': self._prepare_consumers_data(dev_data['consumers']),
            'dev_bindings': self._prepare_bindings_data(dev_data['bindings']),
            'dev_connections': self._prepare_connections_data(dev_data['connections'])
        }
        
        # Combine and export all datasets
        all_datasets = {**prod_datasets, **dev_datasets}
        
        for name, df in all_datasets.items():
            filename = f"{prefix}_{name}.csv"
            df.to_csv(filename, index=False)
            logger.info(f"✓ Created {filename} ({len(df)} rows)")

    def _prepare_comparison_summary(self, comparison_data: Dict) -> pd.DataFrame:
        """Prepare high-level comparison summary"""
        prod_data = comparison_data['production_data']
        dev_data = comparison_data['development_data']
        
        # Get queue names for comparison
        prod_queues = {q['name'] for q in prod_data['queues']}
        dev_queues = {q['name'] for q in dev_data['queues']}
        
        # Get exchange names (excluding default exchanges)
        prod_exchanges = {e['name'] for e in prod_data['exchanges'] if not e['name'].startswith('amq.')}
        dev_exchanges = {e['name'] for e in dev_data['exchanges'] if not e['name'].startswith('amq.')}
        
        # Get consumers by queue
        prod_consumer_queues = {c.get('queue', {}).get('name') for c in prod_data['consumers']}
        dev_consumer_queues = {c.get('queue', {}).get('name') for c in dev_data['consumers']}
        
        summary_data = [
            ['Metric', 'Production', 'Development', 'Difference'],
            ['Generated At', comparison_data['comparison_info']['generated_at'], '', ''],
            ['Production Host', comparison_data['comparison_info']['production']['host'], '', ''],
            ['Development Host', comparison_data['comparison_info']['development']['host'], '', ''],
            ['', '', '', ''],  # Empty row
            ['Total Queues', len(prod_queues), len(dev_queues), len(prod_queues) - len(dev_queues)],
            ['Total Exchanges', len(prod_exchanges), len(dev_exchanges), len(prod_exchanges) - len(dev_exchanges)],
            ['Total Consumers', len(prod_data['consumers']), len(dev_data['consumers']), len(prod_data['consumers']) - len(dev_data['consumers'])],
            ['Total Connections', len(prod_data['connections']), len(dev_data['connections']), len(prod_data['connections']) - len(dev_data['connections'])],
            ['', '', '', ''],  # Empty row
            ['Queues in Prod Only', len(prod_queues - dev_queues), '', ''],
            ['Queues in Dev Only', len(dev_queues - prod_queues), '', ''],
            ['Common Queues', len(prod_queues & dev_queues), '', ''],
            ['', '', '', ''],  # Empty row
            ['Exchanges in Prod Only', len(prod_exchanges - dev_exchanges), '', ''],
            ['Exchanges in Dev Only', len(dev_exchanges - prod_exchanges), '', ''],
            ['Common Exchanges', len(prod_exchanges & dev_exchanges), '', ''],
            ['', '', '', ''],  # Empty row
            ['Queues with Consumers (Prod)', len(prod_consumer_queues), '', ''],
            ['Queues with Consumers (Dev)', len(dev_consumer_queues), '', ''],
            ['Consumer Queue Diff', len(prod_consumer_queues - dev_consumer_queues), '', ''],
        ]
        
        # Add specific queue differences
        if prod_queues - dev_queues:
            summary_data.append(['', '', '', ''])
            summary_data.append(['Queues Only in Production:', '', '', ''])
            for queue in sorted(prod_queues - dev_queues):
                summary_data.append([f"  - {queue}", '', '', ''])
        
        if dev_queues - prod_queues:
            summary_data.append(['', '', '', ''])
            summary_data.append(['Queues Only in Development:', '', '', ''])
            for queue in sorted(dev_queues - prod_queues):
                summary_data.append([f"  - {queue}", '', '', ''])
        
        # Add consumer differences
        consumer_diff = prod_consumer_queues - dev_consumer_queues
        if consumer_diff:
            summary_data.append(['', '', '', ''])
            summary_data.append(['Queues with Consumers Only in Production:', '', '', ''])
            for queue in sorted(consumer_diff):
                summary_data.append([f"  - {queue}", '', '', ''])
        
        return pd.DataFrame(summary_data, columns=['Metric', 'Production', 'Development', 'Difference'])

    def _print_comparison_summary(self, comparison_data: Dict):
        """Print a console summary of the comparison"""
        prod_data = comparison_data['production_data']
        dev_data = comparison_data['development_data']
        
        prod_queues = {q['name'] for q in prod_data['queues']}
        dev_queues = {q['name'] for q in dev_data['queues']}
        
        prod_exchanges = {e['name'] for e in prod_data['exchanges'] if not e['name'].startswith('amq.')}
        dev_exchanges = {e['name'] for e in dev_data['exchanges'] if not e['name'].startswith('amq.')}
        
        prod_consumer_queues = {c.get('queue', {}).get('name') for c in prod_data['consumers'] if c.get('queue', {}).get('name')}
        dev_consumer_queues = {c.get('queue', {}).get('name') for c in dev_data['consumers'] if c.get('queue', {}).get('name')}
        
        print("\n" + "="*80)
        print("RABBITMQ ENVIRONMENT COMPARISON SUMMARY")
        print("="*80)
        print(f"Production:  {comparison_data['comparison_info']['production']['host']}")
        print(f"Development: {comparison_data['comparison_info']['development']['host']}")
        print(f"Generated:   {comparison_data['comparison_info']['generated_at']}")
        print()
        
        # Queues Summary
        print("QUEUES:")
        prod_only_queues = prod_queues - dev_queues
        dev_only_queues = dev_queues - prod_queues
        print(f"  • Production only: {len(prod_only_queues)} queues")
        if prod_only_queues:
            print(f"    {', '.join(sorted(prod_only_queues))}")
        print(f"  • Development only: {len(dev_only_queues)} queues")
        if dev_only_queues:
            print(f"    {', '.join(sorted(dev_only_queues))}")
        print(f"  • Common: {len(prod_queues & dev_queues)} queues")
        print()
        
        # Exchanges Summary
        prod_only_exchanges = prod_exchanges - dev_exchanges
        dev_only_exchanges = dev_exchanges - prod_exchanges
        print("EXCHANGES:")
        print(f"  • Production only: {len(prod_only_exchanges)} exchanges")
        if prod_only_exchanges:
            print(f"    {', '.join(sorted(prod_only_exchanges))}")
        print(f"  • Development only: {len(dev_only_exchanges)} exchanges")
        if dev_only_exchanges:
            print(f"    {', '.join(sorted(dev_only_exchanges))}")
        print(f"  • Common: {len(prod_exchanges & dev_exchanges)} exchanges")
        print()
        
        # Consumers Summary
        consumer_diff = prod_consumer_queues - dev_consumer_queues
        print("CONSUMERS:")
        print(f"  • Total production consumers: {len(prod_data['consumers'])}")
        print(f"  • Total development consumers: {len(dev_data['consumers'])}")
        print(f"  • Queues with consumers (prod only): {len(consumer_diff)}")
        if consumer_diff:
            print(f"    {', '.join(sorted(consumer_diff))}")
        dev_consumer_diff = dev_consumer_queues - prod_consumer_queues
        print(f"  • Queues with consumers (dev only): {len(dev_consumer_diff)}")
        if dev_consumer_diff:
            print(f"    {', '.join(sorted(dev_consumer_diff))}")
        print()


# Helpers
def load_env_files(env_files):
    load_dotenv() # Default .env
    if env_files:
        for env_file in env_files:
            load_dotenv(env_file)